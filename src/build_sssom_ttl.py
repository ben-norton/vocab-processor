"""Generate SSSOM mapping-set TTL files from the SSSOM sheet of each vocabulary xlsx.

Runs alongside build_mappings_ttl.py, which emits the plain SKOS match triples
from the Mappings sheet. This script reads the separate SSSOM sheet and writes
one {slug}_sssom_mappings.ttl file per vocabulary to data/vocabs/, serialised
per the SSSOM RDF format
(https://mapping-commons.github.io/sssom/spec-formats-rdf/):

  [] a sssom:MappingSet ;
      sssom:mapping_set_id    "<vocabularyIRI>/mappings/sssom"^^xsd:anyURI ;
      sssom:license           "<schema:license>"^^xsd:anyURI ;
      sssom:subject_source    <vocabularyIRI> ;
      sssom:mappings          [
          a                         owl:Axiom ;
          owl:annotatedSource       <subject_id> ;
          owl:annotatedProperty     <predicate_id> ;
          owl:annotatedTarget       <object_id> ;
          sssom:mapping_justification semapv:ManualMappingCuration ;
          ...
      ] .

Sheets consumed per workbook:
  Concept Scheme / ConceptScheme  — vocabulary IRI, title and licence
  Namespaces / Namespeces         — extra CURIE prefixes (built-ins win on clash)
  SSSOM                           — one mapping per row, SSSOM slot names as headers

Columns that are not SSSOM slots are ignored and reported (in the console and
the output header). Multivalued slots (author_id, creator_id, ...) split on "|".
Workbooks whose SSSOM sheet is missing or has headers only are skipped.

Run from anywhere (venv active):
    python src/build_sssom_ttl.py [Workbook ...]

Dependencies: openpyxl (required); rdflib (optional, parse check).
"""

from __future__ import annotations

import argparse
import datetime as _dt
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_mappings_ttl import (  # noqa: E402
    OUT_DIR,
    ROOT,
    XLSX_DIR,
    _escape,
    _find_ws,
    iri_slug,
    load_scheme_info,
    normalize_iri,
)

OUT_SUFFIX = "_sssom_mappings.ttl"

# Built-in prefixes. Workbook Namespaces sheets are merged in, but cannot
# override these (several sheets declare skos/vann without the trailing #//).
BUILTIN_PREFIXES: dict[str, str] = {
    "sssom":   "https://w3id.org/sssom/",
    "semapv":  "https://w3id.org/semapv/vocab/",
    "owl":     "http://www.w3.org/2002/07/owl#",
    "rdfs":    "http://www.w3.org/2000/01/rdf-schema#",
    "skos":    "http://www.w3.org/2004/02/skos/core#",
    "xsd":     "http://www.w3.org/2001/XMLSchema#",
    "dcterms": "http://purl.org/dc/terms/",
    "orcid":   "https://orcid.org/",
}

# SSSOM mapping slots → value kind. "ref" = EntityReference (CURIE or IRI →
# IRI node), "str" = plain string, "uri" = xsd:anyURI literal, "double" /
# "date" = typed literals.
MAPPING_SLOTS: dict[str, str] = {
    "subject_label":            "str",
    "subject_category":         "str",
    "subject_type":             "str",
    "subject_source":           "ref",
    "subject_source_version":   "str",
    "subject_match_field":      "ref",
    "subject_preprocessing":    "ref",
    "predicate_label":          "str",
    "predicate_modifier":       "str",
    "object_label":             "str",
    "object_category":          "str",
    "object_type":              "str",
    "object_source":            "ref",
    "object_source_version":    "str",
    "object_match_field":       "ref",
    "object_preprocessing":     "ref",
    "mapping_justification":    "ref",
    "author_id":                "ref",
    "author_label":             "str",
    "reviewer_id":              "ref",
    "reviewer_label":           "str",
    "creator_id":               "ref",
    "creator_label":            "str",
    "license":                  "uri",
    "mapping_provider":         "uri",
    "mapping_cardinality":      "str",
    "mapping_tool":             "str",
    "mapping_tool_version":     "str",
    "mapping_date":             "date",
    "publication_date":         "date",
    "confidence":               "double",
    "similarity_score":         "double",
    "similarity_measure":       "str",
    "curation_rule":            "ref",
    "curation_rule_text":       "str",
    "match_string":             "str",
    "issue_tracker_item":       "ref",
    "see_also":                 "str",
    "other":                    "str",
    "comment":                  "str",
}
MULTIVALUED = {
    "author_id", "author_label", "reviewer_id", "reviewer_label",
    "creator_id", "creator_label", "subject_match_field", "object_match_field",
    "subject_preprocessing", "object_preprocessing", "curation_rule",
    "curation_rule_text", "match_string", "see_also",
}
TRIPLE_SLOTS = ("subject_id", "predicate_id", "object_id")

# ---------------------------------------------------------------------------
# Formatting constants
# ---------------------------------------------------------------------------

INDENT    = "    "
PROP_W    = 26   # "sssom:mapping_set_title" is 23 chars
BN_INDENT = "        "
BN_PROP_W = 34   # "sssom:mapping_justification" is 27 chars

LOCAL_RE = re.compile(r"^[A-Za-z0-9_][A-Za-z0-9_.-]*(?<!\.)$")


class Prefixes:
    """CURIE expansion/compaction, tracking which prefixes the output uses."""

    def __init__(self, extra: list[tuple[str, str]]):
        self.map = dict(BUILTIN_PREFIXES)
        for p, ns in extra:
            if p not in self.map and ns not in self.map.values():
                self.map[p] = ns
        self.used: set[str] = set()

    def expand(self, value: str) -> str:
        """CURIE or IRI → IRI. Raises ValueError for an unknown prefix."""
        v = value.strip()
        if re.match(r"^https?://", v):
            return normalize_iri(v)
        if ":" in v:
            p, local = v.split(":", 1)
            if p in self.map:
                return self.map[p] + local
        raise ValueError(f"unknown prefix in {v!r}")

    def node(self, iri: str) -> str:
        """Render an IRI as a prefixed name when a declared namespace fits."""
        best = max(
            (p for p, ns in self.map.items() if iri.startswith(ns)),
            key=lambda p: len(self.map[p]),
            default=None,
        )
        if best is not None:
            local = iri[len(self.map[best]):]
            if LOCAL_RE.match(local):
                self.used.add(best)
                return f"{best}:{local}"
        return f"<{iri}>"

    def term(self, prefixed: str) -> str:
        self.used.add(prefixed.split(":", 1)[0])
        return prefixed


def lit(s: str) -> str:
    s = str(s).strip()
    if "\n" in s:
        return '"""' + s.replace("\\", "\\\\").replace('"""', '\\"\\"\\"') + '"""'
    return f'"{_escape(s)}"'


def typed(value: str, dtype: str, px: Prefixes) -> str:
    return f'"{_escape(str(value).strip())}"^^{px.term("xsd:" + dtype)}'


def split_multi(raw) -> list[str]:
    return [p.strip() for p in str(raw).split("|") if p.strip()]


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------

def load_license(ws) -> str | None:
    for row in ws.iter_rows(values_only=True):
        if row and row[0] and str(row[0]).strip() in ("schema:license", "sdo:license", "license"):
            if len(row) > 1 and row[1]:
                return str(row[1]).strip()
    return None


def load_namespaces(ws) -> list[tuple[str, str]]:
    pairs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and len(row) > 1 and row[0] and row[1]:
            ns = str(row[1]).strip()
            if ns.startswith(("http://", "https://")) and ns.endswith(("/", "#")):
                pairs.append((str(row[0]).strip(), ns))
    return pairs


def load_sssom(ws) -> tuple[list[str | None], list[tuple]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h not in (None, "") else None for h in rows[0]]
    data = [r for r in rows[1:] if any(v not in (None, "") for v in r)]
    return headers, data


# ---------------------------------------------------------------------------
# Renderer
# ---------------------------------------------------------------------------

def render_mapping(rec: dict, px: Prefixes) -> list[str]:
    def bn(p: str, v: str) -> str:
        return f"{BN_INDENT}{p:<{BN_PROP_W}}{v} ;"

    lines = [bn("a", px.term("owl:Axiom"))]
    lines.append(bn(px.term("owl:annotatedSource"),   px.node(rec["subject_id"])))
    lines.append(bn(px.term("owl:annotatedProperty"), px.node(rec["predicate_id"])))
    lines.append(bn(px.term("owl:annotatedTarget"),   px.node(rec["object_id"])))

    for slot, kind in MAPPING_SLOTS.items():
        vals = rec.get(slot)
        if not vals:
            continue
        if kind == "ref":
            rendered = [px.node(v) for v in vals]
        elif kind == "uri":
            rendered = [typed(v, "anyURI", px) for v in vals]
        elif kind in ("double", "date"):
            rendered = [typed(v, kind, px) for v in vals]
        else:
            rendered = [lit(v) for v in vals]
        lines.append(bn(px.term(f"sssom:{slot}"), " , ".join(rendered)))

    lines[-1] = lines[-1][:-2]  # drop trailing " ;" inside the blank node
    return lines


def render_sssom_ttl(
    records: list[dict],
    px: Prefixes,
    vocab_iri: str,
    vocab_name: str,
    license_iri: str | None,
    source_filename: str,
    notes: list[str],
) -> str:
    def prop(p: str, v: str) -> str:
        return f"{INDENT}{p:<{PROP_W}}{v} ;"

    body = ["[]"]
    body.append(prop("a", px.term("sssom:MappingSet")))
    body.append(prop(px.term("sssom:mapping_set_id"),
                     typed(f"{vocab_iri}/mappings/sssom", "anyURI", px)))
    body.append(prop(px.term("sssom:mapping_set_title"), lit(f"{vocab_name} SSSOM Mappings")))
    if license_iri:
        body.append(prop(px.term("sssom:license"), typed(license_iri, "anyURI", px)))
    body.append(prop(px.term("sssom:subject_source"), f"<{vocab_iri}>"))

    body.append(f"{INDENT}{px.term('sssom:mappings'):<{PROP_W}}[")
    for i, rec in enumerate(records):
        body.extend(render_mapping(rec, px))
        body.append(f"{INDENT}] , [" if i < len(records) - 1 else f"{INDENT}] .")

    today = _dt.date.today().isoformat()
    out = [
        f"# {vocab_name} — SSSOM Mappings",
        f"# Source: {source_filename} (SSSOM sheet)",
        f"# Generated: {today}",
        "#",
        "# SSSOM mapping set, RDF serialisation:",
        "#   https://mapping-commons.github.io/sssom/spec-formats-rdf/",
        f"# Plain SKOS match triples are in {iri_slug(vocab_iri)}-mappings.ttl.",
    ]
    for n in notes:
        out.append(f"# NOTE: {n}")
    out.append("")

    used = sorted(px.used)
    w = max(len(p) + 1 for p in used)
    for p in used:
        out.append(f"@prefix {p + ':':<{w}}  <{px.map[p]}> .")
    out.append("")
    out.extend(body)
    return "\n".join(out) + "\n"


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

def process_workbook(xlsx_path: Path) -> tuple[str, str, int, list[str]] | None:
    """Return (ttl_text, slug, mapping_count, notes) or None if nothing to emit."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    sssom_ws, _ = _find_ws(wb, "SSSOM")
    if sssom_ws is None:
        return None
    headers, rows = load_sssom(sssom_ws)
    if not rows:
        return None

    cs_ws, _ = _find_ws(wb, "Concept Scheme", "ConceptScheme")
    vocab_iri, vocab_name = load_scheme_info(cs_ws) if cs_ws else ("", "")
    if not vocab_iri:
        raise ValueError(f"No vocabularyIRI found in {xlsx_path.name}")
    vocab_name = vocab_name or iri_slug(vocab_iri)
    license_iri = load_license(cs_ws)

    ns_ws, _ = _find_ws(wb, "Namespaces", "Namespeces")
    px = Prefixes(load_namespaces(ns_ws) if ns_ws else [])

    notes: list[str] = []
    missing = [s for s in TRIPLE_SLOTS if s not in headers]
    if missing:
        raise ValueError(f"SSSOM sheet lacks required column(s): {', '.join(missing)}")
    ignored = [h for h in headers if h and h not in MAPPING_SLOTS and h not in TRIPLE_SLOTS]
    if ignored:
        notes.append(f"ignored non-SSSOM column(s): {', '.join(ignored)}")
    if not license_iri:
        notes.append("no schema:license on the Concept Scheme sheet; SSSOM requires a licence")

    records: list[dict] = []
    for n, row in enumerate(rows, start=2):
        cells = {h: row[i] for i, h in enumerate(headers) if h and i < len(row) and row[i] not in (None, "")}
        try:
            rec: dict = {s: px.expand(str(cells.get(s, ""))) for s in TRIPLE_SLOTS if cells.get(s)}
            if len(rec) < 3:
                notes.append(f"row {n}: skipped, missing subject_id/predicate_id/object_id")
                continue
            for slot, kind in MAPPING_SLOTS.items():
                if slot not in cells:
                    continue
                vals = split_multi(cells[slot]) if slot in MULTIVALUED else [str(cells[slot]).strip()]
                if kind == "ref":
                    vals = [px.expand(v) for v in vals]
                elif kind == "date" and hasattr(cells[slot], "date"):
                    vals = [cells[slot].date().isoformat()]
                rec[slot] = vals
        except ValueError as exc:
            notes.append(f"row {n}: skipped, {exc}")
            continue
        records.append(rec)

    if not records:
        return None
    ttl = render_sssom_ttl(records, px, vocab_iri, vocab_name, license_iri, xlsx_path.name, notes)
    return ttl, iri_slug(vocab_iri), len(records), notes


def _parse_check(path: Path) -> str:
    try:
        import rdflib
    except ImportError:
        return ""
    g = rdflib.Graph().parse(path, format="turtle")
    n = len(set(g.subjects(rdflib.RDF.type, rdflib.OWL.Axiom)))
    return f", parsed OK: {n} mappings"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("vocab", nargs="*", help="workbook name(s), e.g. HazardTypes; default all")
    args = ap.parse_args()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wanted = {Path(v).stem for v in args.vocab}
    xlsx_files = sorted(
        p for p in XLSX_DIR.glob("*.xlsx")
        if not p.name.startswith("~$") and p.name != "Agents.xlsx"
        and (not wanted or p.stem in wanted)
    )
    if not xlsx_files:
        print(f"No matching .xlsx files in {XLSX_DIR}")
        sys.exit(1)

    failed = False
    for xlsx in xlsx_files:
        try:
            result = process_workbook(xlsx)
        except Exception as exc:
            print(f"ERROR processing {xlsx.name}: {exc}")
            failed = True
            continue

        if result is None:
            print(f"SKIP {xlsx.name}: no SSSOM sheet, or it has column headers only")
            continue

        ttl, slug, count, notes = result
        out_path = OUT_DIR / f"{slug}{OUT_SUFFIX}"
        out_path.write_text(ttl, encoding="utf-8")
        try:
            check = _parse_check(out_path)
        except Exception as exc:
            check = f", PARSE FAILED: {exc}"
            failed = True
        print(f"wrote {out_path.relative_to(ROOT)}  ({count} mappings, {len(ttl):,} bytes{check})")
        for n in notes:
            print(f"  note: {n}")

    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
