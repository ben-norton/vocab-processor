#!/usr/bin/env python3
"""
Build a classification system's CSV files from its source workbook.

Each classification system lives in sources/classification-systems/<system>/
and is authored in the single *.xlsx workbook in that folder. Every sheet is
written to <sheet name>.csv next to it (UTF-8, CRLF, minimal quoting), so the
workbook is the source of truth and the CSVs are build output. A sheet named
README is documentation only and is skipped.

Derived output and checks:
  - Hierarchy sheets (columns id, parent_id) are checked against the parent
    chain across all hierarchy sheets: ids unique, every parent_id resolves,
    and depth / materialized_path (where present) match the chain. Any
    mismatch aborts the build. path_labels is authored, not checked: it uses
    hand-chosen abbreviations rather than the code column.
  - For the sheet named <stem>_classification, the transitive closure table
    <stem>_closure.csv (ancestor_id, descendant_id, depth; one self-pair per
    node) is generated rather than authored.

Then build the RDF from the CSVs, e.g. for meteorites:
    python src/build_xkos_ap_no.py

Usage:
    python src/build_classification_csvs.py meteorites
    python src/build_classification_csvs.py --all
"""
from __future__ import annotations

import argparse
import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SYSTEMS_DIR = ROOT / "sources" / "classification-systems"
SKIP_SHEETS = {"readme"}


def cell_text(v) -> str:
    """Workbook cell value -> CSV text, undoing Excel's numeric/date typing."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def read_sheet(ws) -> tuple[list[str], list[dict[str, str]]]:
    rows = [[cell_text(c) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(x.strip() for x in r)]  # drop blank rows
    if not rows:
        return [], []
    header = [h.strip() for h in rows[0]]
    while header and not header[-1]:
        header.pop()
    body = [dict(zip(header, r[:len(header)] + [""] * (len(header) - len(r)))) for r in rows[1:]]
    return header, body


def write_csv(path: Path, header: list[str], rows: list[dict[str, str]]) -> None:
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header, lineterminator="\r\n")
        w.writeheader()
        w.writerows(rows)


def check_hierarchy(sheets: dict[str, tuple[list[str], list[dict]]]) -> list[str]:
    """Validate id/parent_id sheets against the combined parent chain."""
    errors: list[str] = []
    by_id: dict[str, dict] = {}
    for name, (header, rows) in sheets.items():
        if not {"id", "parent_id"} <= set(header):
            continue
        for r in rows:
            rid = r["id"].strip()
            if not rid:
                errors.append(f"{name}: row without id: {r}")
            elif rid in by_id:
                errors.append(f"{name}: duplicate id {rid}")
            by_id[rid] = r

    def chain(r: dict) -> list[dict] | None:
        out, seen = [r], {r["id"]}
        while out[-1]["parent_id"].strip():
            p = by_id.get(out[-1]["parent_id"].strip())
            if p is None or p["id"] in seen:
                return None
            out.append(p)
            seen.add(p["id"])
        return out[::-1]

    for name, (header, rows) in sheets.items():
        if not {"id", "parent_id"} <= set(header):
            continue
        for r in rows:
            c = chain(r)
            if c is None:
                errors.append(f"{name}: id {r['id']} has an unresolved or cyclic parent_id {r['parent_id']!r}")
                continue
            if "depth" in header and r["depth"] != str(len(c) - 1):
                errors.append(f"{name}: id {r['id']} depth {r['depth']!r}, parent chain gives {len(c) - 1}")
            mp = "/" + "/".join(x["id"] for x in c) + "/"
            if "materialized_path" in header and r["materialized_path"] != mp:
                errors.append(f"{name}: id {r['id']} materialized_path {r['materialized_path']!r}, parent chain gives {mp!r}")
    return errors


def closure(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """(ancestor, descendant, distance) pairs, root first, one self-pair per node."""
    by_id = {r["id"]: r for r in rows}
    out = []
    for r in rows:
        path = [r["id"]]
        while by_id[path[-1]]["parent_id"]:
            path.append(by_id[path[-1]]["parent_id"])
        for dist in range(len(path) - 1, -1, -1):
            out.append({"ancestor_id": path[dist], "descendant_id": r["id"], "depth": str(dist)})
    return out


def build(system_dir: Path) -> int:
    books = [p for p in system_dir.glob("*.xlsx") if not p.name.startswith("~$")]
    if len(books) != 1:
        print(f"ERROR: expected one .xlsx in {system_dir}, found {[p.name for p in books]}", file=sys.stderr)
        return 1
    book = books[0]
    wb = load_workbook(book, data_only=True, read_only=True)
    sheets = {ws.title: read_sheet(ws) for ws in wb.worksheets if ws.title.lower() not in SKIP_SHEETS}
    wb.close()

    errors = check_hierarchy(sheets)
    if errors:
        print(f"{book.name}: {len(errors)} error(s); no CSVs written", file=sys.stderr)
        for e in errors:
            print("  " + e, file=sys.stderr)
        return 1

    print(f"{book.relative_to(ROOT)}")
    for name, (header, rows) in sheets.items():
        write_csv(system_dir / f"{name}.csv", header, rows)
        print(f"  wrote {name}.csv  ({len(rows)} rows)")
        if name.endswith("_classification") and {"id", "parent_id"} <= set(header):
            stem = name[: -len("_classification")]
            pairs = closure(rows)
            write_csv(system_dir / f"{stem}_closure.csv", ["ancestor_id", "descendant_id", "depth"], pairs)
            print(f"  wrote {stem}_closure.csv  ({len(pairs)} rows, derived)")
    return 0


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("system", nargs="?", help="folder name under sources/classification-systems/")
    ap.add_argument("--all", action="store_true", help="build every classification system")
    args = ap.parse_args()
    if args.all:
        dirs = sorted(p for p in SYSTEMS_DIR.iterdir() if p.is_dir())
    elif args.system:
        dirs = [SYSTEMS_DIR / args.system]
        if not dirs[0].is_dir():
            ap.error(f"no such folder: {dirs[0]}")
    else:
        ap.error("give a system folder name or --all")
    sys.exit(max((build(d) for d in dirs), default=0))


if __name__ == "__main__":
    main()
