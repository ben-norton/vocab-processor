"""Generate SKOS mappings-only TTL files from the Mappings sheet of each vocabulary xlsx.

For each vocabulary xlsx in sources/xlsx/ (excluding Agents.xlsx) writes one
{slug}-mappings.ttl file to data/vocabs/ containing only SKOS match triples.

Sheets consumed per workbook:
  Concept Scheme / ConceptScheme  — vocabulary IRI and name
  Concepts                        — skos:prefLabel lookup per concept IRI
  Mappings / Mapping              — skos:*Match predicates per concept IRI

Output format per concept block:
  <conceptIRI>
      a                    skos:Concept ;
      skos:prefLabel       "..."@en ;
      skos:exactMatch      <...> ;
      skos:closeMatch      <...> ;
      skos:broadMatch      <...> ;
      skos:narrowMatch     <...> ;
      skos:relatedMatch    <...> .

Concepts that have no mappings are silently skipped.

Run from the project root:
    python utils/build_mappings_ttl.py

Dependencies: openpyxl (required).
"""

from __future__ import annotations

import datetime as _dt
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT     = Path(__file__).resolve().parents[1]
XLSX_DIR = ROOT / "sources" / "xlsx"
OUT_DIR  = ROOT / "data" / "vocabs"

# ---------------------------------------------------------------------------
# Column-name alias tables (shared with xlsx-to-ttl-merged.py)
# ---------------------------------------------------------------------------

CS_KEY_ALIASES: dict[str, str] = {
    "dcterms:title":    "title",
    "skos:prefLabel":   "prefLabel",
}

CONCEPT_COL_ALIASES: dict[str, str] = {
    "skos:conceptIRI":   "conceptIRI",
    "skos:prefLabel":    "prefLabel",
    "skos:notation":     "notation",
}

MAPPINGS_COL_ALIASES: dict[str, str] = {
    "skos:relatedMatch": "relatedMatch",
    "skos:closeMatch":   "closeMatch",
    "skos:exactMatch":   "exactMatch",
    "skos:narrowMatch":  "narrowMatch",
    "skos:broadMatch":   "broadMatch",
    "relatedMatchIRI":   "relatedMatch",
    "closeMatchIRI":     "closeMatch",
    "exactMatchIRI":     "exactMatch",
    "narrowMatchIRI":    "narrowMatch",
    "broadMatchIRI":     "broadMatch",
}

MATCH_ORDER = ("exactMatch", "closeMatch", "broadMatch", "narrowMatch", "relatedMatch")

# ---------------------------------------------------------------------------
# Formatting constants
# ---------------------------------------------------------------------------

INDENT = "    "
PROP_W = 20   # "skos:relatedMatch" is 17 chars; 20 gives clean alignment


def _escape(s: str) -> str:
    return str(s).replace("\\", "\\\\").replace('"', '\\"')


def ttl_str(s: str, lang: str = "en") -> str:
    s = str(s).strip()
    if "\n" in s:
        body = s.replace("\\", "\\\\").replace('"""', '\\"\\"\\"')
        return f'"""{body}"""@{lang}'
    return f'"{_escape(s)}"@{lang}'


def _prop(p: str, v: str) -> str:
    return f"{INDENT}{p:<{PROP_W}}{v} ;"


def _multi(p: str, items: list[str]) -> str:
    pad = " " * (len(INDENT) + PROP_W)
    lines = [f"{INDENT}{p:<{PROP_W}}{items[0]}"]
    for item in items[1:]:
        lines.append(f"{pad}{item}")
    result = []
    for i, ln in enumerate(lines):
        result.append(ln + (" ," if i < len(lines) - 1 else " ;"))
    return "\n".join(result)


# ---------------------------------------------------------------------------
# IRI helpers
# ---------------------------------------------------------------------------

def normalize_iri(iri: str) -> str:
    iri = iri.strip()
    if iri.startswith("http://kos.geospecimens.org"):
        iri = "https://" + iri[len("http://"):]
    return iri


def iri_slug(iri: str) -> str:
    return iri.rstrip("/").rsplit("/", 1)[-1]


def split_vals(raw: str | None) -> list[str]:
    if not raw:
        return []
    return [p.strip() for p in re.split(r"[,\n]+", str(raw).strip()) if p.strip()]


# ---------------------------------------------------------------------------
# Sheet helpers
# ---------------------------------------------------------------------------

def _find_ws(wb, *names: str):
    """Return (worksheet, real_name) for the first matching sheet (case-insensitive)."""
    lc = {s.lower(): s for s in wb.sheetnames}
    for n in names:
        real = lc.get(n.lower())
        if real:
            return wb[real], real
    return None, None


def _col_letter_to_idx(col: str) -> int:
    idx = 0
    for c in col.upper():
        idx = idx * 26 + (ord(c) - ord("A") + 1)
    return idx - 1


def _eval_formula(formula: str, row_cells: list, pref_label: str = "") -> str | None:
    m = re.match(r'^="([^"]*)"&([A-Z]+)\d+$', str(formula))
    if not m:
        return None
    base, col_letters = m.group(1), m.group(2)
    col_idx = _col_letter_to_idx(col_letters)
    slug = None
    if col_idx < len(row_cells):
        cell_val = row_cells[col_idx].value
        if cell_val and not str(cell_val).startswith("http"):
            slug = str(cell_val).strip()
    if not slug and pref_label:
        slug = re.sub(r"[^a-z0-9]+", "-", pref_label.lower()).strip("-")
    return (base + slug) if slug else None


def _normalize_concept_col(raw: str) -> str:
    if raw in CONCEPT_COL_ALIASES:
        return CONCEPT_COL_ALIASES[raw]
    m = re.match(r"^[a-zA-Z][\w]*:(.+@[a-z]{2,3}(?:-[A-Za-z0-9]+)?)$", raw)
    if m:
        return m.group(1)
    return raw


def _normalize_mappings_col(raw: str) -> str:
    if raw in MAPPINGS_COL_ALIASES:
        return MAPPINGS_COL_ALIASES[raw]
    local = raw.split(":", 1)[-1] if ":" in raw else raw
    return MAPPINGS_COL_ALIASES.get(local, local)


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------

def load_scheme_info(ws) -> tuple[str, str]:
    """Return (vocabularyIRI, vocab_name) from the Concept Scheme sheet."""
    vocab_iri = ""
    vocab_name = ""
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        raw_key = str(row[0]).strip()
        key = CS_KEY_ALIASES.get(raw_key, raw_key)
        val = str(row[1]).strip() if len(row) > 1 and row[1] not in (None, "") else ""
        if key == "vocabularyIRI" and val:
            vocab_iri = normalize_iri(val)
        if key in ("prefLabel", "title") and val and not vocab_name:
            vocab_name = val
    return vocab_iri, vocab_name


def load_concept_labels(ws) -> dict[str, str]:
    """Return {conceptIRI: prefLabel} from the Concepts sheet.

    Uses data_only=False workbook to resolve formula-based conceptIRI cells.
    """
    rows_cells = list(ws.iter_rows(values_only=False))
    if not rows_cells:
        return {}

    headers: list[str | None] = []
    for c in rows_cells[0]:
        h = c.value
        headers.append(_normalize_concept_col(str(h).strip()) if h and str(h).strip() else None)

    iri_idx  = next((i for i, h in enumerate(headers) if h == "conceptIRI"), None)
    pref_idx = next((i for i, h in enumerate(headers) if h == "prefLabel"),  None)

    labels: dict[str, str] = {}
    for row_cells in rows_cells[1:]:
        row_values = [c.value for c in row_cells]

        concept_iri = None
        if iri_idx is not None and iri_idx < len(row_cells):
            raw = row_cells[iri_idx].value
            if raw and str(raw).startswith("="):
                pref = str(row_values[pref_idx]).strip() if pref_idx is not None and row_values[pref_idx] else ""
                concept_iri = _eval_formula(str(raw), row_cells, pref)
            elif raw:
                concept_iri = str(raw).strip()

        pref_label = None
        if pref_idx is not None and pref_idx < len(row_values) and row_values[pref_idx]:
            pref_label = str(row_values[pref_idx]).strip()

        if concept_iri and pref_label:
            labels[normalize_iri(concept_iri)] = pref_label

    return labels


def load_mappings(ws) -> dict[str, dict[str, list[str]]]:
    """Return {conceptIRI: {matchType: [targetIRI, ...]}} from the Mappings sheet.

    Expects a data_only=True workbook so formula-based cells return cached values.
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}

    headers: list[str | None] = []
    for h in rows[0]:
        if not h or str(h).strip() == "":
            headers.append(None)
        else:
            headers.append(_normalize_mappings_col(str(h).strip()))

    match_cols = set(MATCH_ORDER)
    iri_idx = next((i for i, h in enumerate(headers) if h == "conceptIRI"), 0)

    result: dict[str, dict[str, list[str]]] = {}
    for row in rows[1:]:
        if iri_idx >= len(row) or not row[iri_idx]:
            continue
        concept_iri = normalize_iri(str(row[iri_idx]).strip())
        if not concept_iri.startswith("http"):
            continue
        for ci, h in enumerate(headers):
            if h not in match_cols or ci >= len(row) or not row[ci]:
                continue
            for target in split_vals(str(row[ci])):
                if not target.startswith("http"):
                    continue
                result.setdefault(concept_iri, {}).setdefault(h, [])
                if target not in result[concept_iri][h]:
                    result[concept_iri][h].append(target)
    return result


# ---------------------------------------------------------------------------
# Renderer
# ---------------------------------------------------------------------------

def render_concept_block(
    concept_iri: str,
    pref_label: str | None,
    concept_maps: dict[str, list[str]],
) -> str:
    lines = [f"<{concept_iri}>"]
    lines.append(_prop("a", "skos:Concept"))
    if pref_label:
        lines.append(_prop("skos:prefLabel", ttl_str(pref_label)))

    for match_type in MATCH_ORDER:
        targets = concept_maps.get(match_type, [])
        if len(targets) == 1:
            lines.append(_prop(f"skos:{match_type}", f"<{targets[0]}>"))
        elif targets:
            lines.append(_multi(f"skos:{match_type}", [f"<{t}>" for t in targets]))

    # Replace the trailing " ;" on the last property with " ."
    if len(lines) > 1:
        lines[-1] = lines[-1][:-1] + "."
    else:
        lines.append(".")

    return "\n".join(lines)


def render_mappings_ttl(
    mappings: dict[str, dict[str, list[str]]],
    concept_labels: dict[str, str],
    vocab_name: str,
    source_filename: str,
) -> str:
    today = _dt.date.today().isoformat()
    out: list[str] = []

    out.append(f"# {vocab_name} — SKOS Mappings")
    out.append(f"# Source: {source_filename}")
    out.append(f"# Generated: {today}")
    out.append("#")
    out.append("# This file contains only SKOS match triples.")
    out.append("# The main vocabulary file carries full concept metadata.")
    out.append("")
    out.append("@prefix skos: <http://www.w3.org/2004/02/skos/core#> .")
    out.append("")
    out.append("# " + "=" * 76)
    out.append("# Mappings")
    out.append("# " + "=" * 76)
    out.append("")

    for concept_iri in sorted(mappings.keys()):
        concept_maps = mappings[concept_iri]
        # Skip if no actual match values survived (e.g. all targets filtered out)
        if not any(concept_maps.get(m) for m in MATCH_ORDER):
            continue
        pref_label = concept_labels.get(concept_iri)
        out.append(render_concept_block(concept_iri, pref_label, concept_maps))
        out.append("")

    return "\n".join(out).rstrip() + "\n"


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

def process_workbook(xlsx_path: Path) -> tuple[str, str, int] | None:
    """Return (ttl_text, slug, concept_count) or None if no Mappings sheet."""
    wb_raw  = openpyxl.load_workbook(xlsx_path, data_only=False)
    wb_calc = openpyxl.load_workbook(xlsx_path, data_only=True)

    cs_ws,  _ = _find_ws(wb_calc, "Concept Scheme", "ConceptScheme")
    con_ws, _ = _find_ws(wb_raw,  "Concepts")
    map_ws, _ = _find_ws(wb_calc, "Mappings", "Mapping")

    if map_ws is None:
        return None

    vocab_iri, vocab_name = load_scheme_info(cs_ws) if cs_ws else ("", xlsx_path.stem)
    if not vocab_iri:
        raise ValueError(f"No vocabularyIRI found in {xlsx_path.name}")

    slug = iri_slug(vocab_iri)
    if not vocab_name:
        vocab_name = slug

    concept_labels = load_concept_labels(con_ws) if con_ws else {}
    mappings       = load_mappings(map_ws)

    if not mappings:
        return None

    ttl = render_mappings_ttl(mappings, concept_labels, vocab_name, xlsx_path.name)
    count = sum(
        1 for m in mappings.values()
        if any(m.get(t) for t in MATCH_ORDER)
    )
    return ttl, slug, count


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    xlsx_files = sorted(
        p for p in XLSX_DIR.glob("*.xlsx")
        if not p.name.startswith("~$") and p.name != "Agents.xlsx"
    )
    if not xlsx_files:
        print(f"No .xlsx files found in {XLSX_DIR}")
        return

    for xlsx in xlsx_files:
        try:
            result = process_workbook(xlsx)
        except Exception as exc:
            print(f"ERROR processing {xlsx.name}: {exc}")
            continue

        if result is None:
            print(f"SKIP {xlsx.name}: no Mappings sheet or no data")
            continue

        ttl, slug, count = result
        out_path = OUT_DIR / f"{slug}-mappings.ttl"
        out_path.write_text(ttl, encoding="utf-8")
        print(f"wrote {out_path.relative_to(ROOT)}  ({count} concepts, {len(ttl):,} bytes)")


if __name__ == "__main__":
    main()
