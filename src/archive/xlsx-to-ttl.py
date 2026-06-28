"""Generate SKOS Turtle files from XLSX workbooks in sources/xlsx/.

Reads each *.xlsx in sources/xlsx/ (ignoring Excel lock files), extracts the
Concept Scheme, Concepts, Namespaces/Namespeces, and Agents sheets, and writes
one .ttl file to data/vocabs/ per workbook.

The output filename is derived from the vocabularyIRI slug (not the workbook
filename), so renaming a workbook does not change its output path.

Run from the project root:
    python sources/xlsx-to-ttl.py

Dependencies: openpyxl (required), rdflib (optional, for post-hoc validation).

See sources/xlsx-to-ttl-process.md for full pipeline documentation.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX_DIR = ROOT / "sources" / "xlsx"
OUT_DIR = ROOT / "data" / "vocabs"

# ---------------------------------------------------------------------------
# Prefix definitions
# ---------------------------------------------------------------------------

STANDARD_PREFIXES: list[tuple[str, str]] = [
    ("dcat",       "http://www.w3.org/ns/dcat#"),
    ("dcterms",    "http://purl.org/dc/terms/"),
    ("isoroles",   "http://def.isotc211.org/iso19115-1/2014/CitationAndResponsiblePartyInformation/code/CI_RoleCode/"),
    ("owl",        "http://www.w3.org/2002/07/owl#"),
    ("prov",       "http://www.w3.org/ns/prov#"),
    ("rdf",        "http://www.w3.org/1999/02/22-rdf-syntax-ns#"),
    ("rdfs",       "http://www.w3.org/2000/01/rdf-schema#"),
    ("reg",        "http://purl.org/linked-data/registry#"),
    ("schema",     "https://schema.org/"),
    ("skos",       "http://www.w3.org/2004/02/skos/core#"),
    ("status",     "https://linked.data.gov.au/def/reg-statuses/"),
    ("vann",       "http://purl.org/vocab/vann/"),
    ("vocdermods", "https://linked.data.gov.au/def/vocdermods/"),
    ("xsd",        "http://www.w3.org/2001/XMLSchema#"),
]

STANDARD_NS_SET = {ns for _, ns in STANDARD_PREFIXES}
STANDARD_PREFIX_NAMES = {p for p, _ in STANDARD_PREFIXES}

VOCDERMODS_BASE = "https://linked.data.gov.au/def/vocdermods/"

# ---------------------------------------------------------------------------
# Key / column-name normalization tables
# ---------------------------------------------------------------------------

# Maps both bare and prefix-qualified Concept Scheme row keys to canonical keys.
CS_KEY_ALIASES: dict[str, str] = {
    "dcterms:title":             "title",
    "skos:prefLabel":            "prefLabel",
    "skos:definition":           "definition",
    "dcterms:description":       "description",
    "dcterms:publisher":         "publisher",
    "dcterms:creator":           "creator",
    "dcterms:contributor":       "contributor",
    "dcterms:provenance":        "provenance",
    "dcterms:created":           "created",
    "dcterms:modified":          "modified",
    "dcat:version":              "version",
    "skos:scopeNote":            "scopeNote",
    "vann:usageNote":            "usageNote",
    "skos:historyNote":          "historyNote",
    "skos:note":                 "note",
    "dcat:keyword":              "keywords",
    "schema:keywords":           "keywords",
    "schema:citation":           "citation",
    "prov:wasDerivedFrom":       "derivedFrom",
    "reg:derivationMode":        "derivationMode",
    "regstatus:derivationMode":  "derivationMode",
    "reg:status":                "status",
    "schema:license":            "license",
    "skos:hasTopConcept":        "hasTopConcept",
}

# Maps prefixed Concepts-sheet column headers to canonical keys.
CONCEPT_COL_ALIASES: dict[str, str] = {
    "skos:conceptIRI":    "conceptIRI",
    "skos:prefLabel":     "prefLabel",
    "skos:altLabel":      "altLabel",
    "skos:broader":       "broader",
    "skos:definition":    "definition",
    "dcterms:provenance": "provenance",
    "skos:historyNote":   "historyNote",
    "skos:editorialNote": "editorialNote",
    "skos:changeNote":    "changeNote",
    "vann:usageNote":     "usageNote",
    "skos:inScheme":      "inScheme",
    "reg:status":         "registryStatus",
    "skos:topConceptOf":  "isTopConceptOf",
    "skos:notation":      "notation",
}

# ---------------------------------------------------------------------------
# Text / IRI helpers
# ---------------------------------------------------------------------------

def _escape(s: str) -> str:
    return str(s).replace("\\", "\\\\").replace('"', '\\"')


def ttl_str(s: str, lang: str | None = "en") -> str:
    """Return a Turtle string literal, single-quoted unless multiline."""
    s = str(s).strip()
    tag = f"@{lang}" if lang else ""
    if "\n" in s:
        body = s.replace("\\", "\\\\").replace('"""', '\\"\\"\\"')
        return f'"""{body}"""{tag}'
    return f'"{_escape(s)}"{tag}'


def ttl_long(s: str, lang: str = "en") -> str:
    """Always triple-quoted (used for definitions, descriptions, scope/usage notes)."""
    s = str(s).strip()
    tag = f"@{lang}"
    body = s.replace("\\", "\\\\").replace('"""', '\\"\\"\\"')
    return f'"""{body}"""{tag}'


def ttl_plain(s: str) -> str:
    """Untyped, untagged string literal."""
    return f'"{_escape(str(s).strip())}"'


def split_vals(raw: str | None, sep: str = r"[,\n]+") -> list[str]:
    if not raw:
        return []
    return [p.strip() for p in re.split(sep, str(raw).strip()) if p.strip()]


def parse_date(val) -> str | None:
    if val is None or str(val).strip() in ("", "None"):
        return None
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()
    s = str(val).strip()
    m = re.match(r"^(\d{4})(\d{2})(\d{2})(?:\.\d+)?$", s)
    if m:
        return f"{m[1]}-{m[2]}-{m[3]}"
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", s)
    if m:
        return m[1]
    return None


def iri_slug(iri: str) -> str:
    return iri.rstrip("/").rsplit("/", 1)[-1]


def agent_bnode(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    return f"_:agent_{slug}"


def clean_iri(s: str) -> str:
    """Strip trailing punctuation artifacts (commas, angle brackets) from an IRI."""
    return s.strip().strip("<>").rstrip(",;").strip()


def _normalise_derivation_mode(raw: str) -> str:
    """Return a full IRI for a reg:derivationMode value.

    Handles three forms that appear in spreadsheets:
      - full IRI:  https://linked.data.gov.au/def/vocdermods/subsetting-and-extension
      - reg: IRI:  reg:subsetting-and-extension   (uncommon)
      - bare slug: subsetting-and-extension
    Also strips the erroneous 'regstatus:' prefix that appears in some sheets.
    """
    dm = raw.strip().replace("regstatus:", "").replace("reg:", "").strip()
    if dm.startswith("http"):
        return clean_iri(dm)
    return f"{VOCDERMODS_BASE}{dm}"


# ---------------------------------------------------------------------------
# Turtle formatting
# ---------------------------------------------------------------------------

INDENT = "    "          # 4 spaces for top-level properties
PROP_W = 25              # property column width (value starts at col 29)
BN_INDENT = "        "   # 8 spaces inside attribution blank nodes
BN_PROP_W = 17           # property column width inside attribution blank nodes
INNER_PROP_W = 13        # property column width inside nested blank nodes


def _prop(p: str, v: str, indent: str = INDENT, width: int = PROP_W) -> str:
    return f"{indent}{p:<{width}}{v} ;"


def _multi(
    p: str,
    items: list[str],
    indent: str = INDENT,
    width: int = PROP_W,
) -> str:
    """Property with multiple values, one per line, comma-separated, semicolon at end."""
    pad = " " * (len(indent) + width)
    lines = [f"{indent}{p:<{width}}{items[0]}"]
    for item in items[1:]:
        lines.append(f"{pad}{item}")
    result = []
    for i, ln in enumerate(lines):
        result.append(ln + (" ," if i < len(lines) - 1 else " ;"))
    return "\n".join(result)


# ---------------------------------------------------------------------------
# Sheet loaders
# ---------------------------------------------------------------------------

def _find_ws(wb, *names: str):
    """Return the first worksheet whose name matches any of *names* (case-insensitive)."""
    lc = {s.lower(): s for s in wb.sheetnames}
    for n in names:
        real = lc.get(n.lower())
        if real:
            return wb[real], real
    return None, None


def load_cs_sheet(ws) -> dict[str, list[str]]:
    """Load Concept Scheme sheet: column A is key, columns B+ are values."""
    out: dict[str, list[str]] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        raw_key = str(row[0]).strip()
        key = CS_KEY_ALIASES.get(raw_key, raw_key)
        vals = [str(v).strip() for v in row[1:] if v not in (None, "") and str(v).strip()]
        if vals:
            out.setdefault(key, []).extend(vals)
    return out


def _col_letter_to_idx(col: str) -> int:
    """Convert an Excel column letter (A, B, …, Z, AA, …) to a 0-based index."""
    idx = 0
    for c in col.upper():
        idx = idx * 26 + (ord(c) - ord("A") + 1)
    return idx - 1


def _eval_formula(formula: str, row_cells: list, pref_label: str = "") -> str | None:
    """Resolve a simple Excel concatenation formula like ="base/"&ColRef.

    Falls back to a prefLabel-derived slug when the referenced cell is empty
    or already contains a full IRI (not a usable slug).
    """
    m = re.match(r'^="([^"]*)"&([A-Z]+)\d+$', str(formula))
    if not m:
        return None
    base, col_letters = m.group(1), m.group(2)
    col_idx = _col_letter_to_idx(col_letters)
    slug = None
    if col_idx < len(row_cells):
        cell_val = row_cells[col_idx].value
        if cell_val and not str(cell_val).startswith("http"):
            slug = str(cell_val).strip()
    if not slug and pref_label:
        slug = re.sub(r"[^a-z0-9]+", "-", pref_label.lower()).strip("-")
    return (base + slug) if slug else None


def load_concepts_sheet(ws) -> list[dict]:
    """Load the Concepts sheet.  Handles formula-based conceptIRI cells."""
    rows_cells = list(ws.iter_rows(values_only=False))
    if not rows_cells:
        return []

    headers = []
    for c in rows_cells[0]:
        h = c.value
        if h is None or str(h).strip() == "":
            headers.append(None)
        else:
            s = str(h).strip()
            headers.append(CONCEPT_COL_ALIASES.get(s, s))

    iri_idx  = next((i for i, h in enumerate(headers) if h == "conceptIRI"), None)
    pref_idx = next((i for i, h in enumerate(headers) if h == "prefLabel"),  None)

    records = []
    for row_cells in rows_cells[1:]:
        row_values = [c.value for c in row_cells]
        if not any(v not in (None, "") and str(v).strip() for v in row_values):
            continue

        rec: dict = {}
        for h, v in zip(headers, row_values):
            if h and v not in (None, "") and str(v).strip():
                rec[h] = v if isinstance(v, (int, float, datetime, date)) else str(v).strip()

        # Resolve formula-based conceptIRI cells (e.g. ="base/"&B2)
        if iri_idx is not None:
            raw_iri = row_cells[iri_idx].value
            if raw_iri and str(raw_iri).startswith("="):
                pref_label = (
                    str(row_values[pref_idx]).strip()
                    if pref_idx is not None and row_values[pref_idx]
                    else ""
                )
                resolved = _eval_formula(str(raw_iri), row_cells, pref_label)
                if resolved:
                    rec["conceptIRI"] = resolved

        if rec.get("conceptIRI"):
            records.append(rec)
    return records


def load_namespaces_sheet(ws) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        p  = str(row[0] or "").strip()
        ns = str(row[1] or "").strip()
        if p and ns and (ns.startswith("http://") or ns.startswith("https://")):
            pairs.append((p, ns))
    return pairs


def load_agents_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h not in (None, "") else None for h in rows[0]]
    agents = []
    for row in rows[1:]:
        rec: dict = {}
        for h, v in zip(headers, row):
            if h and v not in (None, "") and str(v).strip():
                rec[h] = str(v).strip()
        if rec:
            agents.append(rec)
    return agents


# ---------------------------------------------------------------------------
# Issue detection
# ---------------------------------------------------------------------------

def detect_issues(
    cs_data: dict,
    concepts: list[dict],
    ns_sheet_name: str | None,
) -> list[str]:
    issues: list[str] = []

    if ns_sheet_name and "namespeces" in ns_sheet_name.lower():
        issues.append('The namespaces sheet is named "Namespeces" (typo for "Namespaces").')

    scheme_iri = str((cs_data.get("vocabularyIRI") or [""])[0]).strip()

    # http:// scheme IRI when concept IRIs use https://
    if scheme_iri.startswith("http://") and any(
        str(r.get("conceptIRI", "")).startswith("https://") for r in concepts
    ):
        issues.append(
            "vocabularyIRI uses http:// scheme. Concept IRIs in this workbook typically "
            "use https://. Consider normalizing to https:// for the scheme IRI."
        )

    # Concept IRI parent(s) vs scheme IRI
    concept_bases: set[str] = set()
    for r in concepts:
        iri = str(r.get("conceptIRI", "")).strip()
        if "/" in iri:
            concept_bases.add(iri.rstrip("/").rsplit("/", 1)[0])
    if concept_bases and concept_bases != {scheme_iri}:
        n = len(concept_bases)
        desc = "\n".join(
            f"       - <{b}> {'(matches scheme)' if b == scheme_iri else '(does NOT match scheme)'}"
            for b in sorted(concept_bases)
        )
        issues.append(
            f"vocabularyIRI is <{scheme_iri}>, but concept IRIs live under {n} distinct "
            f"parent IRI(s):\n{desc}\n"
            "     Transcribed as-written; reconcile in source if this is unintended."
        )

    # skos:inScheme pointing to a different scheme IRI
    bad_inschem: set[str] = set()
    for r in concepts:
        ins = str(r.get("inScheme", "")).strip()
        if ins and ins != scheme_iri:
            bad_inschem.add(ins)
    if bad_inschem:
        desc = "\n".join(f"       - <{s}>" for s in sorted(bad_inschem))
        issues.append(
            f"Some concepts declare skos:inScheme pointing to an IRI other than the "
            f"vocabularyIRI <{scheme_iri}>:\n{desc}\n"
            "     Transcribed as-written; this will cause SKOS scheme membership checks to fail."
        )

    # Truncated hasTopConcept values
    top_raw = cs_data.get("hasTopConcept", [])
    truncated = [v for v in top_raw if str(v).strip().endswith("/")]
    if truncated:
        issues.append(
            f"The scheme-level skos:hasTopConcept cell contains {len(truncated)} truncated "
            "or empty-slug URI(s) (e.g. ending in \"/\" with no concept slug). "
            "skos:hasTopConcept triples were instead derived from concept-level "
            "skos:topConceptOf values, which are complete."
        )

    # Dangling skos:broader references
    concept_iris = {str(r.get("conceptIRI", "")).strip() for r in concepts}
    dangling: list[str] = []
    for r in concepts:
        broader = str(r.get("broader", "")).strip()
        if broader and broader not in concept_iris:
            pref = r.get("prefLabel", r.get("conceptIRI", "?"))
            dangling.append(f"       - '{pref}' → <{broader}>")
    if dangling:
        desc = "\n".join(dangling)
        issues.append(
            f"The following skos:broader references point to concepts NOT present in this "
            f"vocabulary:\n{desc}\n     Transcribed as-written."
        )

    return issues


# ---------------------------------------------------------------------------
# Agent resolution
# ---------------------------------------------------------------------------

def _agent_ref(a: dict) -> str | None:
    """Return the Turtle subject term for an agent: '<URI>' or '_:bnode'."""
    uri  = (a.get("agentURI")  or "").strip()
    name = (a.get("agentName") or "").strip()
    if uri:
        return f"<{uri}>"
    if name:
        return agent_bnode(name)
    return None


def resolve_dc_refs(
    cs_data: dict, agents: list[dict]
) -> tuple[str | None, str | None, str | None]:
    """Return (creator_ref, publisher_ref, contributor_ref) for dcterms: properties."""
    creator_name     = (cs_data.get("creator")     or [None])[0]
    publisher_name   = (cs_data.get("publisher")   or [None])[0]
    contributor_name = (cs_data.get("contributor") or [None])[0]

    creator_ref = publisher_ref = contributor_ref = None

    # Pass 1: resolve by agent role
    for a in agents:
        ref   = _agent_ref(a)
        roles = (a.get("hadRole") or "").lower()
        if not ref:
            continue
        if "creator" in roles and creator_ref is None:
            creator_ref = ref
        if "publisher" in roles and publisher_ref is None:
            publisher_ref = ref
        if "contributor" in roles and contributor_ref is None:
            contributor_ref = ref

    # Pass 2: resolve by name match against CS sheet values
    for a in agents:
        name = (a.get("agentName") or "").strip()
        ref  = _agent_ref(a) or (
            f"<{a['affiliationURI'].strip()}>" if a.get("affiliationURI") else None
        )
        if not ref:
            continue
        if creator_name and name == creator_name and creator_ref is None:
            creator_ref = ref
        if publisher_name and name == publisher_name and publisher_ref is None:
            publisher_ref = ref
        if contributor_name and name == contributor_name and contributor_ref is None:
            contributor_ref = ref

    # Pass 3: fallback — use first agent ref for any still-unresolved slots
    if agents:
        fallback = _agent_ref(agents[0])
        if fallback is None and agents[0].get("affiliationURI"):
            fallback = f"<{agents[0]['affiliationURI'].strip()}>"
        if fallback:
            creator_ref     = creator_ref     or fallback
            publisher_ref   = publisher_ref   or fallback
            contributor_ref = contributor_ref or fallback

    # Final fallback: plain literal from the CS sheet
    if creator_ref is None and creator_name:
        creator_ref = ttl_str(creator_name)
    if publisher_ref is None and publisher_name:
        publisher_ref = ttl_str(publisher_name)
    if contributor_ref is None and contributor_name:
        contributor_ref = ttl_str(contributor_name)

    return creator_ref, publisher_ref, contributor_ref


def build_prov_attributions(agents: list[dict]) -> list[tuple[str, list[str]]]:
    """Return [(agent_ref, [role, ...]), ...], one entry per distinct agent."""
    seen:  dict[str, list[str]] = {}
    order: list[str] = []
    for a in agents:
        ref   = _agent_ref(a)
        roles = [r.strip() for r in (a.get("hadRole") or "").split(",") if r.strip()]
        if not ref:
            continue
        if ref not in seen:
            seen[ref] = []
            order.append(ref)
        seen[ref].extend(roles)
    return [(ref, seen[ref]) for ref in order]


# ---------------------------------------------------------------------------
# Block renderers
# ---------------------------------------------------------------------------

def render_prefix_block(wb_namespaces: list[tuple[str, str]]) -> str:
    w = max(len(p) + 1 for p, _ in STANDARD_PREFIXES)
    lines = ["# Standard prefixes"]
    for p, ns in STANDARD_PREFIXES:
        label = f"{p}:"
        lines.append(f"@prefix {label:<{w}}  <{ns}> .")

    # Omit extras whose namespace or prefix name duplicates a standard one
    extras = [
        (p, ns) for p, ns in wb_namespaces
        if ns not in STANDARD_NS_SET and p not in STANDARD_PREFIX_NAMES
    ]
    if extras:
        we = max(len(p) + 1 for p, _ in extras)
        lines.append("")
        lines.append("# Workbook-specific prefixes (from the Namespaces sheet)")
        for p, ns in extras:
            label = f"{p}:"
            lines.append(f"@prefix {label:<{we}}  <{ns}> .")

    return "\n".join(lines)


def render_concept_scheme(
    cs_data: dict,
    top_concepts: list[str],
    creator_ref: str | None,
    publisher_ref: str | None,
    contributor_ref: str | None,
    prov_attrs: list[tuple[str, list[str]]],
) -> str:
    scheme_iri = str(cs_data["vocabularyIRI"][0]).strip()

    def first(key: str) -> str | None:
        vals = cs_data.get(key)
        return str(vals[0]).strip() if vals else None

    lines: list[str] = [f"<{scheme_iri}>"]

    lines.append(_prop("a", "skos:ConceptScheme"))

    if pf := first("prefLabel"):
        lines.append(_prop("skos:prefLabel", ttl_str(pf)))
    if ti := first("title"):
        lines.append(_prop("dcterms:title", ttl_str(ti)))
    if defn := first("definition"):
        lines.append(_prop("skos:definition", ttl_long(defn)))
    # Always emit description when present, even if identical to definition
    if desc := first("description"):
        lines.append(_prop("dcterms:description", ttl_long(desc)))

    if publisher_ref:
        lines.append(_prop("dcterms:publisher", publisher_ref))
    if creator_ref:
        lines.append(_prop("dcterms:creator", creator_ref))
    if contributor_ref:
        lines.append(_prop("dcterms:contributor", contributor_ref))

    if prov := first("provenance"):
        lines.append(_prop("dcterms:provenance", ttl_str(prov)))

    created = parse_date(first("created"))
    if created:
        lines.append(_prop("dcterms:created", f'"{created}"^^xsd:date'))
    modified = parse_date(first("modified"))
    if modified:
        lines.append(_prop("dcterms:modified", f'"{modified}"^^xsd:date'))

    version_raw = first("version")
    if version_raw:
        ver = parse_date(version_raw) or version_raw
        lines.append(_prop("dcat:version", ttl_plain(ver)))

    if sn := first("scopeNote"):
        lines.append(_prop("skos:scopeNote", ttl_str(sn)))
    if un := first("usageNote"):
        lines.append(_prop("vann:usageNote", ttl_str(un)))
    if hn := first("historyNote"):
        lines.append(_prop("skos:historyNote", ttl_str(hn)))
    if note := first("note"):
        lines.append(_prop("skos:note", ttl_str(note)))

    kws = split_vals(first("keywords"))
    if kws:
        lines.append(_multi("dcat:keyword", [ttl_str(k) for k in kws]))

    if cit := first("citation"):
        lines.append(_prop("schema:citation", ttl_str(cit)))

    # prov:wasDerivedFrom — split by pipe/newline; strip trailing commas from each
    # item (a spreadsheet artifact); emit IRI for http(s) values, @en literal otherwise
    df_raw = first("derivedFrom")
    if df_raw:
        df_items = [
            item.rstrip(",").strip()
            for item in re.split(r"[|\n]+", df_raw)
            if item.strip().rstrip(",")
        ]
        df_rendered = [
            f"<{clean_iri(i)}>" if re.match(r"^https?://", i) else ttl_str(i)
            for i in df_items
            if i
        ]
        if df_rendered:
            lines.append(_multi("prov:wasDerivedFrom", df_rendered))

    if lic := first("license"):
        lines.append(_prop("schema:license", f"<{clean_iri(lic)}>"))

    # reg:derivationMode — the cell may contain a full IRI, a bare slug, or a
    # prefixed value like regstatus:subsetting-and-extension; normalise to IRI
    if dm := first("derivationMode"):
        lines.append(_prop("reg:derivationMode", f"<{_normalise_derivation_mode(dm)}>"))

    if st := first("status"):
        lines.append(_prop("reg:status", f"<{clean_iri(st)}>"))

    # skos:hasTopConcept derived from concept-level isTopConceptOf (authoritative)
    if top_concepts:
        lines.append(_multi("skos:hasTopConcept", [f"<{iri}>" for iri in top_concepts]))

    # prov:qualifiedAttribution blocks
    for agent_ref, roles in prov_attrs:
        role_vals = " , ".join(roles) if roles else "isoroles:unknown"
        lines.extend([
            f"{INDENT}prov:qualifiedAttribution [",
            f"{BN_INDENT}{'a':<{BN_PROP_W}}prov:Attribution ;",
            f"{BN_INDENT}{'prov:agent':<{BN_PROP_W}}{agent_ref} ;",
            f"{BN_INDENT}{'dcat:hadRole':<{BN_PROP_W}}{role_vals} ;",
            f"{INDENT}] ;",
        ])

    lines.append(".")
    return "\n".join(lines)


def render_concept(rec: dict, scheme_iri: str) -> str:
    iri = str(rec["conceptIRI"]).strip()
    lines = [f"<{iri}>"]

    lines.append(_prop("a", "skos:Concept"))

    if pf := str(rec.get("prefLabel", "")).strip():
        lines.append(_prop("skos:prefLabel", ttl_str(pf)))

    alt_raw = str(rec.get("altLabel", "")).strip()
    if alt_raw:
        alts = split_vals(alt_raw)
        if len(alts) == 1:
            lines.append(_prop("skos:altLabel", ttl_str(alts[0])))
        elif alts:
            lines.append(_multi("skos:altLabel", [ttl_str(a) for a in alts]))

    if defn := str(rec.get("definition", "")).strip():
        lines.append(_prop("skos:definition", ttl_long(defn)))

    if broader := str(rec.get("broader", "")).strip():
        lines.append(_prop("skos:broader", f"<{broader}>"))

    if prov := str(rec.get("provenance", "")).strip():
        lines.append(_prop("dcterms:provenance", ttl_str(prov)))

    if hn := str(rec.get("historyNote", "")).strip():
        lines.append(_prop("skos:historyNote", ttl_str(hn)))

    if en := str(rec.get("editorialNote", "")).strip():
        lines.append(_prop("skos:editorialNote", ttl_str(en)))

    if cn := str(rec.get("changeNote", "")).strip():
        lines.append(_prop("skos:changeNote", ttl_str(cn)))

    if un := str(rec.get("usageNote", "")).strip():
        lines.append(_prop("vann:usageNote", ttl_str(un)))

    # skos:inScheme — default to scheme IRI when the column is absent or blank
    inschem = str(rec.get("inScheme", "")).strip() or scheme_iri
    lines.append(_prop("skos:inScheme", f"<{inschem}>"))

    if rs := str(rec.get("registryStatus", "")).strip():
        lines.append(_prop("reg:status", f"<{rs}>"))

    if itco := str(rec.get("isTopConceptOf", "")).strip():
        lines.append(_prop("skos:topConceptOf", f"<{itco}>"))

    if notation := str(rec.get("notation", "")).strip():
        lines.append(_prop("skos:notation", ttl_plain(notation)))

    lines.append(".")
    return "\n".join(lines)


def render_agent_blocks(agents: list[dict]) -> list[str]:
    """Return a list of Turtle blocks (as strings) for all agents."""
    blocks: list[str] = []
    seen_aff: set[str] = set()

    for a in agents:
        uri      = (a.get("agentURI")       or "").strip()
        name     = (a.get("agentName")      or "").strip()
        email    = (a.get("agentEmail")     or "").strip()
        atype    = (a.get("agentType")      or "schema:Thing").strip()
        aff_uri  = (a.get("affiliationURI") or "").strip()
        aff_name = (a.get("affiliationName") or "").strip()

        subject = f"<{uri}>" if uri else (agent_bnode(name) if name else None)
        if subject is None:
            continue

        lines = [subject]
        lines.append(_prop("a", atype))
        if name:
            lines.append(_prop("schema:name", ttl_plain(name)))
        if email:
            lines.append(_prop("schema:email", ttl_plain(email)))

        is_person = "person" in atype.lower()
        is_org    = "organization" in atype.lower() or "organisation" in atype.lower()

        if is_person:
            if aff_uri:
                lines.append(_prop("schema:affiliation", f"<{aff_uri}>"))
        elif is_org:
            if aff_uri:
                lines.append(_prop("schema:affiliation", f"<{aff_uri}>"))
            if aff_name:
                lines.extend([
                    f"{INDENT}schema:parentOrganization [",
                    f"{BN_INDENT}{'a':<{INNER_PROP_W}}schema:Organization ;",
                    f"{BN_INDENT}{'schema:name':<{INNER_PROP_W}}{ttl_plain(aff_name)} ;",
                    f"{INDENT}] ;",
                ])

        lines.append(".")
        blocks.append("\n".join(lines))

        # For person agents: emit a separate block for their affiliation org
        if is_person and aff_uri and aff_uri not in seen_aff:
            seen_aff.add(aff_uri)
            aff_lines = [
                f"<{aff_uri}>",
                _prop("a", "schema:Organization"),
            ]
            if aff_name:
                aff_lines.append(_prop("schema:name", ttl_plain(aff_name)))
            aff_lines.append(".")
            blocks.append("\n".join(aff_lines))

    return blocks


# ---------------------------------------------------------------------------
# Main pipeline
# ---------------------------------------------------------------------------

def render_ttl(xlsx_path: Path) -> tuple[str, str]:
    """Return (ttl_text, scheme_slug) for a given workbook."""
    # data_only=False lets openpyxl see formula strings (e.g. ="base/"&B2 for
    # conceptIRI cells) rather than their cached computed values.
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)

    cs_ws,  _       = _find_ws(wb, "Concept Scheme", "ConceptScheme")
    con_ws, _       = _find_ws(wb, "Concepts")
    ns_ws,  ns_name = _find_ws(wb, "Namespaces", "Namespeces")
    ag_ws,  _       = _find_ws(wb, "Agents")

    cs_data       = load_cs_sheet(cs_ws)       if cs_ws  else {}
    concepts      = load_concepts_sheet(con_ws) if con_ws else []
    wb_namespaces = load_namespaces_sheet(ns_ws) if ns_ws else []
    agents        = load_agents_sheet(ag_ws)    if ag_ws  else []

    scheme_iri = str((cs_data.get("vocabularyIRI") or [""])[0]).strip()
    if not scheme_iri:
        raise ValueError(f"No vocabularyIRI found in {xlsx_path.name}")
    scheme_slug = iri_slug(scheme_iri)

    # Derive top concepts from concept-level isTopConceptOf (authoritative).
    # Deduplicate by concept IRI — NOT by isTopConceptOf value (which is always
    # the same scheme IRI and would otherwise collapse all top concepts to one).
    seen_top:    set[str]  = set()
    top_concepts: list[str] = []
    for r in concepts:
        itco  = str(r.get("isTopConceptOf", "")).strip()
        c_iri = str(r.get("conceptIRI",     "")).strip()
        if itco and c_iri and c_iri not in seen_top:
            seen_top.add(c_iri)
            top_concepts.append(c_iri)

    creator_ref, publisher_ref, contributor_ref = resolve_dc_refs(cs_data, agents)
    prov_attrs = build_prov_attributions(agents)
    issues     = detect_issues(cs_data, concepts, ns_name)

    sheets_used: list[str] = []
    if cs_ws:  sheets_used.append("Concept Scheme")
    if con_ws: sheets_used.append("Concepts")
    if ns_ws:  sheets_used.append(ns_name or "Namespaces")
    if ag_ws:  sheets_used.append("Agents")

    import datetime as _dt
    today = _dt.date.today().isoformat()

    out: list[str] = []

    vocab_name = str((cs_data.get("prefLabel") or cs_data.get("title") or [scheme_slug])[0])
    out.append(f"# {vocab_name} — SKOS ConceptScheme")
    out.append(f"# Generated from {xlsx_path.name}")
    out.append(f"# Source sheets used: {', '.join(sheets_used)}")
    out.append(f"# Generated: {today}")
    if issues:
        out.append("#")
        out.append("# Data issues / transformation notes:")
        for i, issue in enumerate(issues, 1):
            first_line, *rest = issue.splitlines()
            out.append(f"#  {i}. {first_line}")
            for ln in rest:
                out.append(f"#  {ln}")
    out.append("")

    out.append(render_prefix_block(wb_namespaces))
    out.append("")

    out.append("# " + "=" * 76)
    out.append("# ConceptScheme")
    out.append("# " + "=" * 76)
    out.append("")
    out.append(render_concept_scheme(
        cs_data, top_concepts,
        creator_ref, publisher_ref, contributor_ref,
        prov_attrs,
    ))
    out.append("")

    out.append("# " + "=" * 76)
    out.append("# Concepts")
    out.append("# " + "=" * 76)
    out.append("")
    for rec in concepts:
        out.append(render_concept(rec, scheme_iri))
        out.append("")

    if agents:
        agent_blocks = render_agent_blocks(agents)
        if agent_blocks:
            out.append("# " + "=" * 76)
            out.append("# Agents")
            out.append("# " + "=" * 76)
            out.append("")
            for block in agent_blocks:
                out.append(block)
                out.append("")

    return "\n".join(out).rstrip() + "\n", scheme_slug


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_files = sorted(p for p in XLSX_DIR.glob("*.xlsx") if not p.name.startswith("~$"))
    if not xlsx_files:
        print(f"No .xlsx files found in {XLSX_DIR}")
        return

    for xlsx in xlsx_files:
        try:
            ttl, slug = render_ttl(xlsx)
        except Exception as exc:
            print(f"ERROR processing {xlsx.name}: {exc}")
            continue
        out_path = OUT_DIR / f"{slug}.ttl"
        out_path.write_text(ttl, encoding="utf-8")
        print(f"wrote {out_path.relative_to(ROOT)}  ({len(ttl):,} bytes)")

    # Optional post-hoc validation with rdflib
    try:
        import rdflib
        SKOS = rdflib.Namespace("http://www.w3.org/2004/02/skos/core#")
        print("\nValidating generated files…")
        for ttl_path in sorted(OUT_DIR.glob("*.ttl")):
            g = rdflib.Graph().parse(ttl_path, format="turtle")
            schemes  = list(g.subjects(rdflib.RDF.type, SKOS.ConceptScheme))
            concepts = list(g.subjects(rdflib.RDF.type, SKOS.Concept))
            ok = len(schemes) == 1 and concepts
            print(
                f"  [{'OK' if ok else 'WARN'}] {ttl_path.name}: "
                f"{len(schemes)} scheme(s), {len(concepts)} concept(s)"
            )
    except ImportError:
        pass


if __name__ == "__main__":
    main()
