"""Generate prez-lite-conformant SKOS TTL files from sources/xlsx workbooks.

Reads each vocabulary *.xlsx in sources/xlsx/ (skipping Agents.xlsx) and writes
a corresponding .ttl file to data/vocabs/.  Sheets consumed per workbook:

  Concept Scheme  — key-value metadata for the skos:ConceptScheme
  Concepts        — one row per concept
  Mappings        — SKOS match properties per concept
  Agents          — agent descriptions + role attribution

Column headers that carry namespace prefixes (e.g. ``skos:prefLabel``) are
normalised to their local name (``prefLabel``) so the same handler works
regardless of whether a file uses prefixed or bare column names.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX_DIR = ROOT / "sources" / "xlsx"
OUT_DIR = ROOT / "data" / "vocabs"

PREFIXES: list[tuple[str, str]] = [
    ("dcat",     "http://www.w3.org/ns/dcat#"),
    ("dcterms",  "http://purl.org/dc/terms/"),
    ("isoroles", "http://def.isotc211.org/iso19115-1/2014/"
                 "CitationAndResponsiblePartyInformation/code/CI_RoleCode/"),
    ("owl",      "http://www.w3.org/2002/07/owl#"),
    ("prov",     "http://www.w3.org/ns/prov#"),
    ("rdfs",     "http://www.w3.org/2000/01/rdf-schema#"),
    ("reg",      "http://purl.org/linked-data/registry#"),
    ("schema",   "https://schema.org/"),
    ("skos",     "http://www.w3.org/2004/02/skos/core#"),
    ("vann",     "http://purl.org/vocab/vann/"),
    ("xsd",      "http://www.w3.org/2001/XMLSchema#"),
]

# After stripping the namespace prefix, apply these canonical aliases so that
# files using different local-name conventions resolve to the same key.
_KEY_ALIASES: dict[str, str] = {
    "keyword":         "keywords",       # dcat:keyword  → keywords
    "wasDerivedFrom":  "derivedFrom",    # prov:wasDerivedFrom → derivedFrom
    "relatedMatchIRI": "relatedMatch",
    "closeMatchIRI":   "closeMatch",
    "exactMatchIRI":   "exactMatch",
    "narrowMatchIRI":  "narrowMatch",
    "broadMatchIRI":   "broadMatch",
    "rdfType":         "type",           # bare "rdfType" key used in some files
}


# ── helpers ───────────────────────────────────────────────────────────────────

def normalize_key(key: str) -> str:
    """Strip namespace prefix then apply canonical aliases."""
    local = key.split(":", 1)[-1] if ":" in key else key
    return _KEY_ALIASES.get(local, local)


def normalize_iri(iri: str) -> str:
    iri = iri.strip()
    if iri.startswith("http://kos.geospecimens.org"):
        iri = "https://" + iri[len("http://"):]
    return iri


def parse_date(val) -> str | None:
    if val is None or val == "":
        return None
    if isinstance(val, (datetime, date)):
        d = val.date() if isinstance(val, datetime) else val
        return d.isoformat()
    s = str(val).strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})(\d{2})(\d{2})(?:\.0+)?$", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", s)
    if m:
        return m.group(1)
    return None


def iri_slug(iri: str) -> str:
    return iri.rstrip("/").rsplit("/", 1)[-1]


def slugify(name: str) -> str:
    """Convert a display name to a blank-node-safe slug."""
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def escape_string(s: str) -> str:
    return str(s).replace("\\", "\\\\").replace('"', '\\"')


def ttl_string(s: str, lang: str | None = "en") -> str:
    s = str(s).strip()
    if "\n" in s:
        body = s.replace("\\", "\\\\").replace('"""', '\\"\\"\\"')
        return f'"""{body}"""' + (f"@{lang}" if lang else "")
    return f'"{escape_string(s)}"' + (f"@{lang}" if lang else "")


def ttl_plain(s: str) -> str:
    return f'"{escape_string(str(s).strip())}"'


def split_list(val: str | None) -> list[str]:
    if val is None:
        return []
    s = str(val).strip()
    if not s:
        return []
    return [p.strip() for p in re.split(r"[,\n]+", s) if p.strip()]


# ── xlsx loaders ─────────────────────────────────────────────────────────────

def load_sheet_kv(ws) -> dict[str, list]:
    """Return a dict mapping each normalised key (col A) to its non-empty values."""
    out: dict[str, list] = {}
    for row in ws.iter_rows(values_only=True):
        row = list(row)
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        key = normalize_key(str(row[0]).strip())
        vals = [v for v in row[1:] if v not in (None, "") and str(v).strip() != ""]
        if key in out:
            out[key].extend(vals)
        else:
            out[key] = list(vals)
    return out


def load_concepts(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    # Normalise headers; language-tag variants like "prefLabel@de" are kept as-is
    # after stripping the namespace prefix.
    raw_headers = [str(h).strip() if h is not None else None for h in rows[0]]
    headers: list[str | None] = []
    for h in raw_headers:
        if not h:
            headers.append(None)
        elif "@" in h and ":" in h:
            # e.g. "skos:prefLabel@de" → "prefLabel@de"
            prefix_part, rest = h.split(":", 1)
            headers.append(rest)
        else:
            headers.append(normalize_key(h))

    records = []
    for row in rows[1:]:
        if not any(c not in (None, "") for c in row):
            continue
        rec: dict = {}
        for h, v in zip(headers, row):
            if h and v not in (None, ""):
                val = str(v).strip() if not isinstance(v, (int, float, datetime, date)) else v
                if val != "":
                    rec[h] = val
        if rec.get("conceptIRI"):
            records.append(rec)
    return records


def load_agents(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_key(str(h).strip()) if h is not None else None for h in rows[0]]
    agents = []
    for row in rows[1:]:
        rec: dict = {}
        for h, v in zip(headers, row):
            if h and v not in (None, ""):
                rec[h] = str(v).strip()
        if rec:
            agents.append(rec)
    return agents


def load_mappings(ws) -> dict[str, dict[str, list[str]]]:
    """Return {conceptIRI: {matchType: [targetIRI, ...]}} from a Mappings sheet."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    headers = [normalize_key(str(h).strip()) if h else None for h in rows[0]]
    match_cols = {"relatedMatch", "closeMatch", "exactMatch", "narrowMatch", "broadMatch"}
    concept_col = 0  # conceptIRI is always first
    result: dict[str, dict[str, list[str]]] = {}
    for row in rows[1:]:
        concept_iri = row[concept_col] if row[concept_col] else None
        if not concept_iri:
            continue
        concept_iri = str(concept_iri).strip()
        for ci, h in enumerate(headers):
            if h not in match_cols:
                continue
            val = row[ci] if ci < len(row) else None
            if not val:
                continue
            for target in split_list(str(val)):
                if not target.startswith("http"):
                    continue
                result.setdefault(concept_iri, {}).setdefault(h, [])
                if target not in result[concept_iri][h]:
                    result[concept_iri][h].append(target)
    return result


# ── TTL renderers ─────────────────────────────────────────────────────────────

def render_agent(a: dict) -> tuple[str | None, str | None]:
    """Render an agent description block.

    Returns (ttl_block, subject) where subject is the IRI or blank-node label
    used as the RDF subject (needed for ConceptScheme attribution).
    """
    uri = a.get("agentURI") or a.get("iri")
    name = a.get("agentName") or a.get("name") or ""
    atype = a.get("agentType", "schema:Thing")

    if uri:
        subject = f"<{uri.strip()}>"
    elif name:
        subject = f"_:agent_{slugify(name)}"
    else:
        return None, None

    lines = [subject]
    lines.append(f"    a {atype} ;")
    if name:
        lines.append(f"    schema:name {ttl_plain(name)} ;")
    if a.get("agentEmail"):
        lines.append(f"    schema:email {ttl_plain(a['agentEmail'])} ;")
    # affiliationURI covers the old column name; memberOfURI the new one
    aff = a.get("affiliationURI") or a.get("memberOfURI")
    if aff:
        lines.append(f"    schema:affiliation <{aff.strip()}> ;")
    lines.append(".")
    return "\n".join(lines), subject


def render_concept(
    rec: dict,
    scheme_iri: str,
    mappings: dict[str, dict[str, list[str]]],
) -> tuple[str, str, bool]:
    """Render one concept block.

    Returns (ttl_block, concept_iri, is_top_concept).
    """
    concept_iri = normalize_iri(str(rec["conceptIRI"]))

    lines = [f"<{concept_iri}>"]
    lines.append("    a skos:Concept ;")

    if rec.get("prefLabel"):
        lines.append(f"    skos:prefLabel {ttl_string(rec['prefLabel'])} ;")

    # Language-specific prefLabels (e.g. "prefLabel@de")
    for key, val in rec.items():
        m = re.fullmatch(r"prefLabel@([a-z]{2,3}(?:-[A-Za-z0-9]+)?)", key)
        if m:
            lang = m.group(1)
            lines.append(f"    skos:prefLabel {ttl_string(val, lang=lang)} ;")

    if rec.get("altLabel"):
        labels = split_list(rec["altLabel"])
        if len(labels) == 1:
            lines.append(f"    skos:altLabel {ttl_string(labels[0])} ;")
        elif labels:
            joined = " ,\n        ".join(ttl_string(lb) for lb in labels)
            lines.append(f"    skos:altLabel\n        {joined} ;")

    if rec.get("notation"):
        lines.append(f"    skos:notation {ttl_plain(rec['notation'])} ;")

    if rec.get("definition"):
        lines.append(f"    skos:definition {ttl_string(rec['definition'])} ;")
    if rec.get("provenance"):
        lines.append(f"    dcterms:provenance {ttl_string(rec['provenance'])} ;")
    if rec.get("historyNote"):
        lines.append(f"    skos:historyNote {ttl_string(rec['historyNote'])} ;")
    if rec.get("editorialNote"):
        lines.append(f"    skos:editorialNote {ttl_string(rec['editorialNote'])} ;")
    if rec.get("changeNote"):
        lines.append(f"    skos:changeNote {ttl_string(rec['changeNote'])} ;")
    if rec.get("usageNote"):
        lines.append(f"    vann:usageNote {ttl_string(rec['usageNote'])} ;")

    if rec.get("broader"):
        broader_iri = normalize_iri(str(rec["broader"]))
        lines.append(f"    skos:broader <{broader_iri}> ;")

    lines.append(f"    skos:inScheme <{scheme_iri}> ;")

    if rec.get("status"):
        lines.append(f"    reg:status <{rec['status'].strip()}> ;")

    is_top = bool(rec.get("topConceptOf"))
    if is_top:
        lines.append(f"    skos:topConceptOf <{scheme_iri}> ;")

    # SKOS match properties from Mappings sheet
    concept_maps = mappings.get(concept_iri, {})
    for match_type in ("relatedMatch", "closeMatch", "exactMatch", "narrowMatch", "broadMatch"):
        for target in concept_maps.get(match_type, []):
            pred = f"skos:{match_type}"
            lines.append(f"    {pred} <{target}> ;")

    lines.append(".")
    return "\n".join(lines), concept_iri, is_top


# ── main render ───────────────────────────────────────────────────────────────

def render_ttl(xlsx_path: Path) -> tuple[str, str]:
    """Read an xlsx file and return (ttl_text, scheme_slug)."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet_map = {s.lower(): s for s in wb.sheetnames}

    # Sheet name resolution (tolerate spelling variants)
    cs_sheet      = sheet_map.get("concept scheme") or sheet_map.get("conceptscheme")
    concepts_sheet = sheet_map.get("concepts")
    agents_sheet  = sheet_map.get("agents")
    mappings_sheet = sheet_map.get("mappings") or sheet_map.get("mapping")

    cs_data   = load_sheet_kv(wb[cs_sheet]) if cs_sheet else {}
    concepts  = load_concepts(wb[concepts_sheet]) if concepts_sheet else []
    agents    = load_agents(wb[agents_sheet]) if agents_sheet else []
    mappings  = load_mappings(wb[mappings_sheet]) if mappings_sheet else {}

    scheme_iri = normalize_iri(str(cs_data["vocabularyIRI"][0]))
    scheme_slug = iri_slug(scheme_iri)

    created  = parse_date(cs_data.get("created",  [None])[0])
    modified = parse_date(cs_data.get("modified", [None])[0])
    version_raw = cs_data.get("version", [None])[0]
    version = parse_date(version_raw) or (str(version_raw) if version_raw else None)

    pref_label  = (cs_data.get("prefLabel")   or cs_data.get("title")      or [scheme_slug])[0]
    title       = (cs_data.get("title")       or cs_data.get("prefLabel")  or [pref_label])[0]
    definition  = (cs_data.get("definition")  or cs_data.get("description") or [""])[0]
    description = (cs_data.get("description") or [""])[0]

    # Rendered concept blocks
    concept_blocks: list[str] = []
    top_concepts: list[str] = []
    for rec in concepts:
        block, concept_iri, is_top = render_concept(rec, scheme_iri, mappings)
        concept_blocks.append(block)
        if is_top:
            top_concepts.append(concept_iri)

    seen: set[str] = set()
    top_concepts = [c for c in top_concepts if not (c in seen or seen.add(c))]  # type: ignore[func-returns-value]

    # Resolve agent subjects and build role attribution data
    agent_blocks: list[str] = []
    agent_subjects: list[tuple[str, str]] = []  # (subject, hadRole)
    for a in agents:
        block, subject = render_agent(a)
        if block:
            agent_blocks.append(block)
        if subject:
            role = a.get("hadRole", "")
            if role:
                agent_subjects.append((subject, role))

    # Resolve creator / publisher / contributor IRIs from agents
    creator_iri = publisher_iri = None
    contributor_iris: list[str] = []
    creator_name  = (cs_data.get("creator",  [None])[0] or "")
    publisher_name = (cs_data.get("publisher", [None])[0] or "")

    for a in agents:
        roles = (a.get("hadRole") or "").lower()
        name  = a.get("agentName") or a.get("name") or ""
        uri   = a.get("agentURI") or a.get("iri") or a.get("affiliationURI")
        if not uri:
            continue
        uri = uri.strip()
        if ("editor" in roles or "creator" in roles or "author" in roles) and creator_iri is None:
            creator_iri = uri
        if "publisher" in roles and publisher_iri is None:
            publisher_iri = uri
        if creator_name and name == creator_name and creator_iri is None:
            creator_iri = uri
        if publisher_name and name == publisher_name and publisher_iri is None:
            publisher_iri = uri
        if "contributor" in roles or "custodian" in roles:
            if uri not in contributor_iris:
                contributor_iris.append(uri)

    # Fallback: use any affiliationURI
    if creator_iri is None or publisher_iri is None:
        for a in agents:
            aff = a.get("affiliationURI")
            if aff:
                creator_iri   = creator_iri   or aff.strip()
                publisher_iri = publisher_iri or aff.strip()
                break

    # ── Build TTL ─────────────────────────────────────────────────────────────
    out: list[str] = []

    # Prefix declarations
    out.append(f"PREFIX :   <{scheme_iri}/>")
    out.append(f"PREFIX cs: <{scheme_iri}>")
    for p, ns in PREFIXES:
        out.append(f"PREFIX {p}: <{ns}>")
    out.append("")

    # ConceptScheme block
    cs_lines: list[str] = ["cs:"]
    cs_lines.append("    a")
    cs_lines.append("        owl:Ontology ,")
    cs_lines.append("        skos:ConceptScheme ;")
    cs_lines.append(f"    dcterms:title {ttl_string(title)} ;")
    cs_lines.append(f"    skos:prefLabel {ttl_string(pref_label)} ;")
    if definition:
        cs_lines.append(f"    skos:definition {ttl_string(definition)} ;")
    if description and description != definition:
        cs_lines.append(f"    dcterms:description {ttl_string(description)} ;")
    if cs_data.get("scopeNote"):
        cs_lines.append(f"    skos:scopeNote {ttl_string(cs_data['scopeNote'][0])} ;")
    if cs_data.get("usageNote"):
        cs_lines.append(f"    vann:usageNote {ttl_string(cs_data['usageNote'][0])} ;")
    if cs_data.get("historyNote"):
        cs_lines.append(f"    skos:historyNote {ttl_string(cs_data['historyNote'][0])} ;")
    if cs_data.get("note"):
        cs_lines.append(f"    skos:note {ttl_string(cs_data['note'][0])} ;")
    if cs_data.get("provenance"):
        cs_lines.append(f"    dcterms:provenance {ttl_string(cs_data['provenance'][0])} ;")

    # Keywords
    kw_raw = cs_data.get("keywords", [None])[0]
    if kw_raw:
        kws = split_list(kw_raw)
        if kws:
            joined = " ,\n        ".join(ttl_string(k) for k in kws)
            cs_lines.append(f"    dcat:keyword\n        {joined} ;")

    if cs_data.get("citation"):
        cs_lines.append(f"    schema:citation {ttl_string(cs_data['citation'][0])} ;")

    # Derived-from
    df_raw = (cs_data.get("derivedFrom") or [None])[0]
    if df_raw:
        for d in split_list(df_raw):
            if d.startswith("http"):
                cs_lines.append(f"    prov:wasDerivedFrom <{d}> ;")

    # Derivation mode
    dm = (cs_data.get("derivationMode") or [None])[0]
    if dm and dm.startswith("http"):
        cs_lines.append(f"    reg:derivationMode <{dm.strip()}> ;")

    if cs_data.get("status"):
        cs_lines.append(f"    reg:status <{cs_data['status'][0].strip()}> ;")
    if cs_data.get("license"):
        cs_lines.append(f"    schema:license <{cs_data['license'][0].strip()}> ;")

    # Agent attribution on ConceptScheme
    if creator_iri:
        cs_lines.append(f"    dcterms:creator <{creator_iri}> ;")
    elif creator_name:
        cs_lines.append(f"    dcterms:creator {ttl_plain(creator_name)} ;")
    if publisher_iri:
        cs_lines.append(f"    dcterms:publisher <{publisher_iri}> ;")
    elif publisher_name:
        cs_lines.append(f"    dcterms:publisher {ttl_plain(publisher_name)} ;")
    for c_iri in contributor_iris:
        cs_lines.append(f"    dcterms:contributor <{c_iri}> ;")

    if created:
        cs_lines.append(f'    dcterms:created "{created}"^^xsd:date ;')
    if modified:
        cs_lines.append(f'    dcterms:modified "{modified}"^^xsd:date ;')
    if version:
        cs_lines.append(f"    dcat:version {ttl_plain(version)} ;")

    # Top concepts
    if top_concepts:
        rendered_tops = [f"<{iri}>" for iri in top_concepts]
        joined = " ,\n        ".join(rendered_tops)
        cs_lines.append(f"    skos:hasTopConcept\n        {joined} ;")

    # prov:qualifiedAttribution per agent with a role
    for subject, had_role in agent_subjects:
        roles = [r.strip() for r in had_role.split(",") if r.strip()]
        if not roles:
            continue
        roles_ttl = " , ".join(roles)
        cs_lines.append("    prov:qualifiedAttribution [")
        cs_lines.append("        a prov:Attribution ;")
        cs_lines.append(f"        prov:agent {subject} ;")
        cs_lines.append(f"        dcat:hadRole {roles_ttl} ;")
        cs_lines.append("    ] ;")

    cs_lines.append(".")
    out.append("\n".join(cs_lines))
    out.append("")

    # Agent description blocks
    for block in agent_blocks:
        out.append(block)
        out.append("")

    # Concept blocks
    for block in concept_blocks:
        out.append(block)
        out.append("")

    return "\n".join(out).rstrip() + "\n", scheme_slug


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for xlsx in sorted(XLSX_DIR.glob("*.xlsx")):
        if xlsx.stem.startswith("~") or xlsx.name == "Agents.xlsx":
            continue
        try:
            ttl, slug = render_ttl(xlsx)
        except Exception as exc:
            print(f"ERROR {xlsx.name}: {exc}")
            continue
        out_path = OUT_DIR / f"{slug}.ttl"
        out_path.write_text(ttl, encoding="utf-8")
        print(f"wrote {out_path.relative_to(ROOT)}  ({len(ttl):,} bytes)")


if __name__ == "__main__":
    main()
